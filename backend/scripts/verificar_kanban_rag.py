"""Verificacao no VPS: leitor de documentos + biblioteca do iaMAC, kanban com justificativa, KPIs."""
import os
import io
import json
import sqlite3
import sys
import time
import urllib.error
import urllib.request
import uuid

API = os.environ.get("VISIONDESK_API", "http://localhost:8200/api")
ADMIN_PIN, OPER_PIN, SUP_PIN = sys.argv[1], sys.argv[2], sys.argv[3]
REL = {"ok": [], "falhas": []}


def http(method, path, body=None, token=None, timeout=240, raw=None, ctype="application/json"):
    data = raw if raw is not None else (json.dumps(body).encode() if body is not None else None)
    req = urllib.request.Request(API + path, data=data, method=method)
    if data is not None:
        req.add_header("Content-Type", ctype)
    if token:
        req.add_header("X-Auth-Token", token)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, json.loads(r.read().decode() or "{}")
    except urllib.error.HTTPError as e:
        try:
            return e.code, json.loads(e.read().decode() or "{}")
        except Exception:
            return e.code, {}
    except Exception as e:  # noqa: BLE001
        return 0, {"erro": str(e)[:120]}


def multipart(files):
    b = uuid.uuid4().hex
    out = io.BytesIO()
    for name, data in files:
        out.write(f"--{b}\r\nContent-Disposition: form-data; name=\"files\"; filename=\"{name}\"\r\nContent-Type: application/octet-stream\r\n\r\n".encode())
        out.write(data); out.write(b"\r\n")
    out.write(f"--{b}--\r\n".encode())
    return out.getvalue(), f"multipart/form-data; boundary={b}"


def ok(nome, cond, det=""):
    (REL["ok"] if cond else REL["falhas"]).append(nome + (f" — {det}" if det else ""))
    print(("  [ok]   " if cond else "  [FALHA] ") + nome + (f" — {det}" if det else ""), flush=True)


def login(oid, pin):
    s, d = http("POST", "/auth/login/pin", {"operator_id": oid, "pin": pin})
    return d.get("token")


ADM = login("op_admin_demo", ADMIN_PIN); OPER = login("op_operations", OPER_PIN); SUP = login("op_supply", SUP_PIN)
ok("logins", bool(ADM and OPER and SUP))

print("\n== 1. leitor de documentos: inventario e pergunta 'quais arquivos'")
s, d = http("GET", "/iamac/documents")
ok("inventario de documentos", s == 200 and d["indexados"] > 40, f"{d.get('indexados')} indexados de {d.get('total')} · formatos {len(d.get('supported', []))}")
s, a = http("POST", "/iamac/ask", {"question": "quais arquivos e documentos você tem na base?"})
fontes = [c["source"] for c in a.get("citations", [])]
ok("pergunta sobre arquivos cita o Inventário", "Inventário de documentos" in fontes, str(fontes[:3]))

print("\n== 2. biblioteca: upload multiformato, leitura e remocao")
import docx  # noqa: E402
import openpyxl  # noqa: E402
dd = docx.Document(); dd.add_heading("Procedimento de Bloqueio e Etiquetagem LOTO", 1)
dd.add_paragraph("Antes de intervir no painel, o operador deve desenergizar, bloquear a chave seccionadora e etiquetar com o cadeado pessoal. Teste de ausência de tensão obrigatório.")
bd = io.BytesIO(); dd.save(bd)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "KPIs"; ws.append(["indicador", "meta", "real"]); ws.append(["OEE", 85, 93.1]); ws.append(["OTIF", 95, 100])
bx = io.BytesIO(); wb.save(bx)
arqs = [("loto_procedimento.docx", bd.getvalue()), ("kpis_setembro.xlsx", bx.getvalue()),
        ("norma_iec62443_resumo.txt", "IEC 62443 define zonas e conduítes; o nível de seguranca SL-T deve ser definido por zona. Segmentacao da rede OT e defesa em profundidade.".encode()),
        ("catalogo.html", "<html><body><h1>Catalogo de Sensores</h1><p>Sensor indutivo M12 PNP, alcance 4mm, IP67.</p><script>x=1</script></body></html>".encode()),
        ("imagem.bmp", b"BM\x00\x00\x00\x00")]
raw, ct = multipart(arqs)
s, up = http("POST", "/iamac/library/upload", raw=raw, ctype=ct, token=ADM)
res = {r["file"]: r for r in up.get("results", [])}
ok("docx lido", res.get("loto_procedimento.docx", {}).get("lido") is True, f"{res.get('loto_procedimento.docx', {}).get('chunks')} trechos")
ok("xlsx lido", res.get("kpis_setembro.xlsx", {}).get("lido") is True, f"{res.get('kpis_setembro.xlsx', {}).get('chunks')} trechos")
ok("txt lido", res.get("norma_iec62443_resumo.txt", {}).get("lido") is True)
ok("html lido (sem script)", res.get("catalogo.html", {}).get("lido") is True)
ok("formato nao suportado e recusado com motivo", res.get("imagem.bmp", {}).get("ok") is False, res.get("imagem.bmp", {}).get("message", ""))
s, a = http("POST", "/iamac/ask", {"question": "qual o procedimento de bloqueio e etiquetagem antes de intervir no painel?"})
ok("RAG responde com o documento enviado (docx)", any("loto" in c["source"] for c in a.get("citations", [])), str([c["source"] for c in a.get("citations", [])][:2]))
s, a = http("POST", "/iamac/ask", {"question": "qual a meta e o real de OEE na planilha de KPIs?"})
ok("RAG responde com a planilha (xlsx)", any("kpis_setembro" in c["source"] for c in a.get("citations", [])), (a.get("answer") or "")[:90].replace("\n", " "))
s, d = http("POST", "/iamac/library/upload", raw=raw, ctype=ct, token=None)
ok("upload sem sessao e recusado (403)", s == 403)
for n in ["loto_procedimento.docx", "kpis_setembro.xlsx", "norma_iec62443_resumo.txt", "catalogo.html"]:
    http("DELETE", f"/iamac/library/{n}", token=ADM)
s, lb = http("GET", "/iamac/library")
ok("biblioteca limpa apos remocao", s == 200 and len(lb.get("files", [])) == 0, f"{len(lb.get('files', []))} arquivos")

print("\n== 3. kanban: movimentos com justificativa")
s, b = http("GET", "/vps/board?entity=OP&segment=PANEL")
cards = b.get("cards", [])
ok("board com phase_entered_ts e tempo na fase", all("phase_entered_ts" in c for c in cards), f"{len(cards)} cards")
livre = next((c for c in cards if c["state"] == "IN_PROCESS" and c["pendings"] == 0 and c["phase"] not in ("DONE", "BILLING")), None)
bloq = next((c for c in cards if c["state"] == "EXCEPTION" and c["pendings"] > 0), None)
done = next((c for c in cards if c["phase"] == "DONE"), None)
ok("ha OP livre, OP bloqueada e OP concluida para testar", bool(livre and bloq and done),
   f"livre={livre and livre['natural_key']} bloq={bloq and bloq['natural_key']} done={done and done['natural_key']}")
if livre:
    tag = livre["tag_id"]; phases = livre["phases"]; i = phases.index(livre["phase"])
    # supply nao e dono de OP → 403
    s, d = http("POST", f"/vps/thread/{tag}/move", {"to_phase": phases[min(i + 1, len(phases) - 1)], "reason": "teste de permissao supply"}, SUP)
    ok(f"op_supply mover OP → {s} (esp 403)", s == 403, (d.get("detail") or {}).get("message", "")[:70])
    # sem justificativa → 409
    s, d = http("POST", f"/vps/thread/{tag}/move", {"to_phase": phases[min(i + 1, len(phases) - 1)], "reason": "x"}, OPER)
    ok(f"sem justificativa → {s} (esp 409)", s == 409, (d.get("detail") or {}).get("message", "")[:70])
    # avancar 1 com justificativa
    s, d = http("POST", f"/vps/thread/{tag}/advance", {"reason": "material conferido, liberado pela qualidade — teste automatizado"}, OPER)
    ok(f"avancar 1 fase com justificativa → {s}", s == 200 and d.get("ok"), f"{livre['phase']} → {d.get('phase')} [{d.get('state')}]")
    s, th = http("GET", f"/vps/thread/{tag}")
    ult = th["timeline"][-1] if th.get("timeline") else {}
    pay = ult.get("payload") or {}
    ok("evento carrega reason e moved_by=op_operations", pay.get("reason", "").startswith("material") and pay.get("moved_by") == "op_operations" and ult.get("operator_id") == "op_operations",
       f"{ult.get('action')} por {ult.get('operator_id')}")
    # pular 2 fases → ocorrencia PHASE_SKIPPED
    s, b2 = http("GET", "/vps/board?entity=OP&segment=PANEL"); c2 = next(c for c in b2["cards"] if c["tag_id"] == tag)
    j = phases.index(c2["phase"]) if c2["phase"] in phases else 0
    alvo = phases[min(j + 2, len(phases) - 1)]
    s, d = http("POST", f"/vps/thread/{tag}/move", {"to_phase": alvo, "reason": "urgencia do cliente: montagem e teste feitos em bancada externa"}, OPER)
    ok(f"pular 2 fases → {s} com ocorrencia PHASE_SKIPPED", s == 200 and len(d.get("occurrences", [])) == 1, f"{c2['phase']} → {d.get('to')} fired={d.get('fired')}")
    # voltar 1 fase → REWORK
    s, b3 = http("GET", "/vps/board?entity=OP&segment=PANEL"); c3 = next(c for c in b3["cards"] if c["tag_id"] == tag)
    k = phases.index(c3["phase"]) if c3["phase"] in phases else 1
    volta = phases[max(k - 1, 0)]
    s, d = http("POST", f"/vps/thread/{tag}/move", {"to_phase": volta, "reason": "retrabalho: solda irregular detectada na inspecao"}, OPER)
    ok(f"voltar 1 fase → {s} REWORK com ocorrencia", s == 200 and "REWORK" in d.get("fired", []) and d.get("phase") == volta, f"{c3['phase']} → {d.get('to')} [{d.get('state')}]")
    s, th = http("GET", f"/vps/thread/{tag}")
    occs = [o["kind"] for o in th.get("occurrences", [])]
    ok("ocorrencias PHASE_SKIPPED e REWORK abertas no Fio", "PHASE_SKIPPED" in occs and "REWORK" in occs, str(occs))
if bloq:
    s, d = http("POST", f"/vps/thread/{bloq['tag_id']}/move", {"to_phase": "DONE", "reason": "tentativa de mover card bloqueado"}, ADM)
    ok(f"OP bloqueada por pendencia → {s} (esp 409)", s == 409, (d.get("detail") or {}).get("message", "")[:70])
if done:
    s, d = http("POST", f"/vps/thread/{done['tag_id']}/move", {"to_phase": "PICKING", "reason": "tentativa de reabrir concluida pelo kanban"}, ADM)
    ok(f"OP concluida nao se move → {s} (esp 409)", s == 409, (d.get("detail") or {}).get("message", "")[:70])

print("\n== 4. KPIs")
s, k = http("GET", "/vps/kpis")
ok("reconciliacao por instancia (nao por evento)", s == 200 and isinstance(k.get("CLD_Reconciliation_Rate"), (int, float)), f"recon={k.get('CLD_Reconciliation_Rate')}% leadtime={k.get('CLD_Thread_LeadTime_h')}h otif={k.get('CLD_OTIF')}% ader={k.get('CLD_Schedule_Adherence')}% custo={k.get('CLD_Cost_of_Pendencies')}")

print("\n== 5. trilha")
c = sqlite3.connect("/var/lib/visiondesk/visiondesk.db")
mv = c.execute("SELECT COUNT(*) FROM audit_event WHERE action='CARD_MOVED'").fetchone()[0]
lib = c.execute("SELECT COUNT(*) FROM audit_event WHERE action IN ('LIBRARY_UPLOAD','LIBRARY_DELETE')").fetchone()[0]
ok("CARD_MOVED auditado", mv >= 2, str(mv)); ok("biblioteca auditada", lib >= 8, str(lib))
print("\n== RESUMO"); print(json.dumps({"ok": len(REL["ok"]), "falhas": REL["falhas"]}, ensure_ascii=False, indent=1))
